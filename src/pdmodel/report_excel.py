"""Formatted Excel monitoring report (openpyxl)."""
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
HEADER_FONT = Font(bold=True, color="FFFFFF")
RAG_FILL = {"Green": PatternFill("solid", fgColor="C6EFCE"),
            "Amber": PatternFill("solid", fgColor="FFEB9C"),
            "Red": PatternFill("solid", fgColor="FFC7CE")}
RAG_FONT = {"Green": Font(color="006100", bold=True), "Amber": Font(color="9C5700", bold=True),
            "Red": Font(color="9C0006", bold=True)}
THIN = Border(bottom=Side(style="thin", color="D9D9D9"))


def _write_table(ws, df, start_row=1, formats=None):
    formats = formats or {}
    for j, col in enumerate(df.columns, 1):
        c = ws.cell(row=start_row, column=j, value=col)
        c.fill, c.font = HEADER_FILL, HEADER_FONT
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for i, row in enumerate(df.itertuples(index=False), start_row + 1):
        for j, (col, val) in enumerate(zip(df.columns, row), 1):
            if isinstance(val, float) and pd.isna(val):
                val = None
            c = ws.cell(row=i, column=j, value=val)
            c.border = THIN
            if col in formats:
                c.number_format = formats[col]
            if isinstance(val, str) and val in RAG_FILL:
                c.fill, c.font = RAG_FILL[val], RAG_FONT[val]
                c.alignment = Alignment(horizontal="center")
    for j, col in enumerate(df.columns, 1):
        width = max(len(str(col)), *(len(f"{v:.4f}" if isinstance(v, float) else str(v)) for v in df[col])) + 2
        ws.column_dimensions[get_column_letter(j)].width = min(max(width, 10), 40)
    ws.freeze_panes = ws.cell(row=start_row + 1, column=1)
    return start_row + len(df) + 2


def build_report(path, monitoring, coefficients, points, thresholds_text, dev_metrics):
    wb = Workbook()

    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = "Corporate PD Model - Quarterly Monitoring Report (simulated quarters)"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = (f"Development (hold-out) Gini: {dev_metrics['gini']:.3f} | AUC: {dev_metrics['auc']:.3f} | "
                f"KS: {dev_metrics['ks']:.3f}")
    s = monitoring["summary"].copy()
    s = s[["quarter", "risk_tilt", "feature_noise", "dr_uplift", "n_firms", "defaults", "observed_dr", "mean_pd",
           "psi", "psi_rag", "max_csi", "max_csi_variable", "gini", "gini_relative_drop", "gini_rag",
           "binomial_p", "calibration_rag", "overall_rag"]]
    row = _write_table(ws, s, start_row=4, formats={
        "observed_dr": "0.0%", "mean_pd": "0.0%", "psi": "0.000", "max_csi": "0.000", "gini": "0.000",
        "gini_relative_drop": "0.0%", "binomial_p": "0.0000", "risk_tilt": "0.00", "feature_noise": "0.00",
        "dr_uplift": "0.00"})
    ws.cell(row=row, column=1, value="RAG thresholds").font = Font(bold=True)
    for k, line in enumerate(thresholds_text, 1):
        ws.cell(row=row + k, column=1, value=line)

    _write_table(wb.create_sheet("PSI detail"), monitoring["psi_detail"],
                 formats={"dev_share": "0.0%", "current_share": "0.0%", "psi_contribution": "0.0000"})
    csi_fmt = {c: "0.000" for c in monitoring["csi"].columns if c != "quarter"}
    _write_table(wb.create_sheet("CSI by variable"), monitoring["csi"], formats=csi_fmt)
    g = monitoring["grades"].drop(columns=["dr_ci_low", "dr_ci_high"])
    _write_table(wb.create_sheet("Calibration by grade"), g, formats={
        "pd_lower": "0.00%", "pd_upper": "0.00%", "mean_pd": "0.00%", "share": "0.0%",
        "observed_dr": "0.00%", "binomial_p": "0.0000"})
    _write_table(wb.create_sheet("Model coefficients"), coefficients, formats={
        c: "0.0000" for c in ["coefficient", "std_error", "z", "p_value", "ci_lower", "ci_upper"]})
    _write_table(wb.create_sheet("Scorecard points"), points,
                 formats={"default_rate": "0.0%", "woe": "0.000"})
    for sheet in wb.worksheets:             # print-friendly: landscape, one page wide
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
    wb.save(path)
